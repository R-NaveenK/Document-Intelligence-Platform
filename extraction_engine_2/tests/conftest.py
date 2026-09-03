"""Pytest fixtures and test document generators for COR extraction engine."""
import io
from pathlib import Path
import tempfile
import pytest
from PIL import Image, ImageDraw
import fitz
import docx


@pytest.fixture(scope="session")
def test_output_dir():
    """Temporary output directory for test extractions."""
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@pytest.fixture(scope="session")
def sample_invoice_data():
    return {
        "title": "TAX INVOICE",
        "vendor": "TECHSOLUTIONS PVT LTD",
        "address": "123, 4th Floor, Prestige Tech Park, Bengaluru - 560103",
        "gstin": "29AABCT1234Q1Z5",
        "invoice_no": "TSPL/INV/2026-27/0897",
        "invoice_date": "02-09-2026",
        "customer": "ABC INDUSTRIES LTD",
        "item_1": "Dell Latitude 5420 Laptop",
        "item_1_qty": "5",
        "item_1_price": "58500",
        "item_1_total": "292500",
        "item_2": "Logitech MX Master 3S Mouse",
        "item_2_qty": "10",
        "item_2_price": "9450",
        "item_2_total": "94500",
        "subtotal": "387000",
        "tax_rate": "18%",
        "tax_amount": "69660",
        "grand_total": "Rs. 456,660",
    }


def create_invoice_image(invoice_data: dict, width: int = 800, height: int = 700) -> Image.Image:
    """Render a synthetic invoice as a PIL Image."""
    img = Image.new("RGB", (width, height), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    lines = [
        (invoice_data["title"], 30, 20),
        (invoice_data["vendor"], 30, 60),
        (invoice_data["address"], 30, 85),
        (f"GSTIN: {invoice_data['gstin']}", 30, 110),
        (f"Invoice No: {invoice_data['invoice_no']}", 30, 140),
        (f"Invoice Date: {invoice_data['invoice_date']}", 30, 165),
        (f"BILL TO: {invoice_data['customer']}", 30, 200),
        ("ITEMS:", 30, 240),
        (f"{invoice_data['item_1']} | Qty: {invoice_data['item_1_qty']} | Rate: {invoice_data['item_1_price']} | Amt: {invoice_data['item_1_total']}", 30, 270),
        (f"{invoice_data['item_2']} | Qty: {invoice_data['item_2_qty']} | Rate: {invoice_data['item_2_price']} | Amt: {invoice_data['item_2_total']}", 30, 300),
        (f"Subtotal: {invoice_data['subtotal']}", 30, 340),
        (f"GST ({invoice_data['tax_rate']}): {invoice_data['tax_amount']}", 30, 365),
        (f"Grand Total: {invoice_data['grand_total']}", 30, 400),
    ]

    for text, x, y in lines:
        draw.text((x, y), text, fill=(0, 0, 0))

    return img


@pytest.fixture
def digital_pdf_file(tmp_path, sample_invoice_data):
    """Create a digital PDF with searchable embedded text."""
    pdf_path = tmp_path / "digital_invoice.pdf"
    doc = fitz.open()
    page = doc.new_page()

    lines = [
        sample_invoice_data["title"],
        sample_invoice_data["vendor"],
        sample_invoice_data["address"],
        f"GSTIN: {sample_invoice_data['gstin']}",
        f"Invoice No: {sample_invoice_data['invoice_no']}",
        f"Invoice Date: {sample_invoice_data['invoice_date']}",
        f"Bill To: {sample_invoice_data['customer']}",
        f"Item: {sample_invoice_data['item_1']} Qty: {sample_invoice_data['item_1_qty']} Total: {sample_invoice_data['item_1_total']}",
        f"Grand Total: {sample_invoice_data['grand_total']}",
    ]
    y = 50
    for line in lines:
        page.insert_text((50, y), line, fontsize=11)
        y += 25

    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def scanned_pdf_file(tmp_path, sample_invoice_data):
    """Create a scanned PDF where pages only contain image pixels (no embedded text)."""
    pdf_path = tmp_path / "scanned_invoice.pdf"
    img = create_invoice_image(sample_invoice_data)
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    img_bytes.seek(0)

    doc = fitz.open()
    page = doc.new_page(width=img.width, height=img.height)
    rect = fitz.Rect(0, 0, img.width, img.height)
    page.insert_image(rect, stream=img_bytes.read())

    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def mixed_pdf_file(tmp_path, sample_invoice_data):
    """Create a 2-page PDF: Page 1 is digital text, Page 2 is a scanned image."""
    pdf_path = tmp_path / "mixed_invoice.pdf"
    doc = fitz.open()

    # Page 1: Digital text
    page1 = doc.new_page()
    page1.insert_text((50, 50), "PAGE 1: SUMMARY\nTAX INVOICE DETAILS\nVendor: TECHSOLUTIONS PVT LTD\nTotal: Rs. 456,660")

    # Page 2: Image scan
    img = create_invoice_image(sample_invoice_data)
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    img_bytes.seek(0)

    page2 = doc.new_page(width=img.width, height=img.height)
    page2.insert_image(fitz.Rect(0, 0, img.width, img.height), stream=img_bytes.read())

    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def multipage_pdf_file(tmp_path):
    """Create a multi-page digital PDF with 3 pages."""
    pdf_path = tmp_path / "multipage_doc.pdf"
    doc = fitz.open()
    for i in range(1, 4):
        page = doc.new_page()
        page.insert_text((50, 50), f"Section {i} Content: Detailed report paragraph number {i} with metrics and totals.")
    doc.save(pdf_path)
    doc.close()
    return pdf_path


@pytest.fixture
def png_image_file(tmp_path, sample_invoice_data):
    """Create a PNG image file containing invoice text."""
    p = tmp_path / "invoice_sample.png"
    img = create_invoice_image(sample_invoice_data)
    img.save(p)
    return p


@pytest.fixture
def jpg_image_file(tmp_path, sample_invoice_data):
    """Create a JPG image file containing invoice text."""
    p = tmp_path / "invoice_sample.jpg"
    img = create_invoice_image(sample_invoice_data)
    img.save(p, "JPEG")
    return p


@pytest.fixture
def docx_file(tmp_path, sample_invoice_data):
    """Create a DOCX file containing paragraphs, headers, and an items table."""
    docx_path = tmp_path / "invoice_doc.docx"
    doc = docx.Document()

    # Heading & Paragraphs
    doc.add_heading("TAX INVOICE", level=1)
    doc.add_paragraph(f"Vendor: {sample_invoice_data['vendor']}")
    doc.add_paragraph(f"GSTIN: {sample_invoice_data['gstin']}")
    doc.add_paragraph(f"Invoice No: {sample_invoice_data['invoice_no']}")
    doc.add_paragraph(f"Customer: {sample_invoice_data['customer']}")

    # Table
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item Description"
    hdr_cells[1].text = "Quantity"
    hdr_cells[2].text = "Unit Price"
    hdr_cells[3].text = "Line Total"

    row1 = table.add_row().cells
    row1[0].text = sample_invoice_data["item_1"]
    row1[1].text = sample_invoice_data["item_1_qty"]
    row1[2].text = sample_invoice_data["item_1_price"]
    row1[3].text = sample_invoice_data["item_1_total"]

    row2 = table.add_row().cells
    row2[0].text = sample_invoice_data["item_2"]
    row2[1].text = sample_invoice_data["item_2_qty"]
    row2[2].text = sample_invoice_data["item_2_price"]
    row2[3].text = sample_invoice_data["item_2_total"]

    doc.add_paragraph(f"Grand Total: {sample_invoice_data['grand_total']}")
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def txt_file(tmp_path, sample_invoice_data):
    """Create a TXT file containing structured-like invoice lines."""
    p = tmp_path / "invoice_plain.txt"
    content = f"""{sample_invoice_data['title']}

Vendor: {sample_invoice_data['vendor']}
GSTIN: {sample_invoice_data['gstin']}
Invoice No: {sample_invoice_data['invoice_no']}
Date: {sample_invoice_data['invoice_date']}

Bill To: {sample_invoice_data['customer']}

Item 1: {sample_invoice_data['item_1']}
Quantity: {sample_invoice_data['item_1_qty']}
Price: {sample_invoice_data['item_1_price']}
Total: {sample_invoice_data['item_1_total']}

Grand Total: {sample_invoice_data['grand_total']}
"""
    p.write_text(content, encoding="utf-8")
    return p
