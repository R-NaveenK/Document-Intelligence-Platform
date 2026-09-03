"""Tests for DOCX extraction handler."""
import io
from pathlib import Path
from PIL import Image, ImageDraw
import docx
from cor_engine.extractor import CorExtractor
from cor_engine.models import ExtractionStatus


def test_docx_extraction(docx_file, test_output_dir):
    extractor = CorExtractor(output_dir=test_output_dir)
    result = extractor.extract(docx_file, document_id="invoice_docx")

    assert result.status == ExtractionStatus.SUCCESS
    assert result.file_type == "DOCX"
    assert result.page_count >= 1
    assert result.character_count > 50

    content = result.raw_text
    # Heading & paragraphs
    assert "TAX INVOICE" in content
    assert "TECHSOLUTIONS PVT LTD" in content
    assert "29AABCT1234Q1Z5" in content

    # Table items extracted as unstructured lines
    assert "Item Description" in content
    assert "Dell Latitude 5420 Laptop" in content
    assert "58500" in content
    assert "Grand Total: Rs. 456,660" in content


def test_docx_embedded_image_ocr(tmp_path, test_output_dir):
    """Test DOCX containing an embedded image with readable text."""
    p = tmp_path / "embedded_img_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("DOCUMENT WITH EMBEDDED STAMP:")

    # Create a stamp image
    img = Image.new("RGB", (300, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 40), "APPROVED FOR PAYMENT", fill=(0, 0, 0))
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    img_bytes.seek(0)

    doc.add_picture(img_bytes)
    doc.save(p)

    extractor = CorExtractor(output_dir=test_output_dir)
    result = extractor.extract(p, document_id="docx_embedded")

    assert result.status == ExtractionStatus.SUCCESS
    assert "DOCUMENT WITH EMBEDDED STAMP" in result.raw_text
    # OCR of the embedded image should recover the text
    assert "APPROVED" in result.raw_text or "PAYMENT" in result.raw_text
