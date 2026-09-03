"""DOCX Document Handler extracting paragraphs, headings, tables, and embedded images."""
import logging
from pathlib import Path
from typing import List

import docx
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from cor_engine.core.exceptions import DocxProcessingError
from cor_engine.core.result import PageExtraction, RawExtraction
from cor_engine.handlers.base import BaseDocumentHandler
from cor_engine.ocr.paddle_engine import PaddleOCREngine

logger = logging.getLogger("cor_engine.handlers.docx")


class DocxHandler(BaseDocumentHandler):
    """Handler for Microsoft Word .docx files preserving document flow and embedded images."""

    def __init__(self):
        self._ocr_engine = PaddleOCREngine.get_instance()

    def process(self, file_path: Path) -> RawExtraction:
        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            logger.error("[COR] Failed to open DOCX '%s': %s", file_path.name, exc)
            raise DocxProcessingError(f"Failed to open DOCX file: {str(exc)}") from exc

        extracted_blocks: List[str] = []

        try:
            # 1. Extract headers where accessible
            for sec in doc.sections:
                try:
                    if sec.header and sec.header.paragraphs:
                        header_text = "\n".join(p.text.strip() for p in sec.header.paragraphs if p.text.strip())
                        if header_text and header_text not in extracted_blocks:
                            extracted_blocks.append(header_text)
                except Exception as h_exc:
                    logger.debug("[COR] Could not extract header: %s", h_exc)

            # 2. Iterate document body elements in document order
            for child in doc.element.body.iterchildren():
                if isinstance(child, CT_P):
                    p = Paragraph(child, doc)
                    text = p.text.strip()
                    if text:
                        extracted_blocks.append(text)
                elif isinstance(child, CT_Tbl):
                    tbl = Table(child, doc)
                    # Extract table content row-by-row, cell-by-cell without JSON structuring
                    table_lines: List[str] = []
                    for row in tbl.rows:
                        row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        # Deduplicate adjacent identical cells from merged cells
                        deduped_cells: List[str] = []
                        for c in row_cells:
                            if not deduped_cells or c != deduped_cells[-1]:
                                deduped_cells.append(c)
                        if deduped_cells:
                            table_lines.append("  |  ".join(deduped_cells))

                    if table_lines:
                        extracted_blocks.append("\n".join(table_lines))

            # 3. Extract footers where accessible
            for sec in doc.sections:
                try:
                    if sec.footer and sec.footer.paragraphs:
                        footer_text = "\n".join(p.text.strip() for p in sec.footer.paragraphs if p.text.strip())
                        if footer_text and footer_text not in extracted_blocks:
                            extracted_blocks.append(footer_text)
                except Exception as f_exc:
                    logger.debug("[COR] Could not extract footer: %s", f_exc)

            # 4. Extract and OCR embedded images containing text
            embedded_img_count = 0
            for rel in doc.part.rels.values():
                if "image" in rel.target_ref:
                    try:
                        img_bytes = rel.target_part.blob
                        ocr_text, conf = self._ocr_engine.process_image(img_bytes, apply_preprocessing=True)
                        if ocr_text.strip():
                            embedded_img_count += 1
                            header = (
                                "----------------------------------------\n"
                                f"EMBEDDED IMAGE {embedded_img_count}\n"
                                "----------------------------------------"
                            )
                            extracted_blocks.append(f"{header}\n{ocr_text.strip()}")
                            logger.info("[COR] Extracted text from embedded image %d in DOCX", embedded_img_count)
                    except Exception as img_exc:
                        logger.debug("[COR] Embedded image skipped or non-text: %s", img_exc)

        except DocxProcessingError:
            raise
        except Exception as exc:
            logger.error("[COR] Error processing DOCX content: %s", exc)
            raise DocxProcessingError(f"DOCX extraction failed: {str(exc)}") from exc

        full_text = "\n\n".join(extracted_blocks).strip()
        pages = [PageExtraction(page_number=1, text=full_text, is_ocr=False, confidence=1.0)]
        return RawExtraction(pages=pages, has_page_numbers=False)

    def extract(self, file_path: Path) -> RawExtraction:
        """Extract method conforming to BaseDocumentHandler."""
        return self.process(file_path)
